from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "docs" / "03_Arquitectura_y_Base_de_Datos"
OUT = OUT_DIR / "04_Especificacion_Tecnica_API_REST_Swagger.docx"
OUT_COPY = ROOT / "docs" / "04_Especificacion_Tecnica_API_REST_Swagger.docx"


def shade_run(run, hex_color):
    run.font.color.rgb = RGBColor.from_string(hex_color)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = str(value)
    doc.add_paragraph()
    return table


def code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    title = doc.add_paragraph()
    run = title.add_run("OficiosYa")
    run.bold = True
    run.font.size = Pt(14)
    shade_run(run, "1F3A44")

    h = doc.add_paragraph()
    r = h.add_run("Especificación técnica API REST / Swagger")
    r.bold = True
    r.font.size = Pt(22)
    shade_run(r, "1F3A44")

    meta = doc.add_paragraph()
    meta.add_run(
        "HU-03 · Lógica de recuperación de contraseña\n"
        "IDC-21 · HU-05 / HU-07 · Perfil y disponibilidad\n"
        "Base URL: http://localhost:3000/api\n"
        "Swagger UI: http://localhost:3000/api/docs\n"
        "Versión 1.1 · 28/09/2026"
    )

    doc.add_heading("1. Alcance de esta entrega", 1)
    doc.add_paragraph(
        "Esta especificación documenta el flujo backend de recuperación de acceso (HU-03): "
        "generación de un token temporal seguro, código de 6 dígitos y simulación del envío "
        "del correo (no hay SMTP real). También conserva GET/PUT /worker/profile y "
        "PUT /worker/availability (IDC-21)."
    )

    doc.add_heading("2. Recuperación de acceso (HU-03)", 1)
    doc.add_paragraph(
        "El usuario pide restablecer la clave. El API busca el correo en usuario. Si la cuenta "
        "existe y está operable, emite un desafío de un solo uso y simula el envío. La respuesta "
        "pública es la misma si el correo no está registrado, para no enumerar cuentas."
    )
    add_table(
        doc,
        ["Paso", "Detalle"],
        [
            ["1. Solicitud", "POST /api/auth/forgot-password { correo }"],
            ["2. Token", "32 bytes aleatorios (crypto.randomBytes) en base64url"],
            ["3. Código", "6 dígitos (crypto.randomInt)"],
            ["4. Almacenamiento", "Solo hashes SHA-256 o HMAC-SHA256 (RESET_TOKEN_SECRET)"],
            ["5. TTL", "RESET_TOKEN_TTL_MINUTES (por defecto 15)"],
            ["6. Envío", "MailSimulatorService escribe en la consola del API"],
            ["7. Uso", "POST /api/auth/reset-password con token o correo+código"],
            ["8. Consumo", "Un solo uso. 5 intentos fallidos de código invalidan el desafío"],
            ["9. Clave nueva", "bcrypt costo 12 en usuario.password_hash + Auth de Supabase"],
        ],
    )

    doc.add_heading("3. POST /auth/forgot-password", 1)
    add_table(
        doc,
        ["Campo", "Valor"],
        [
            ["Método / ruta", "POST /api/auth/forgot-password"],
            ["Auth", "Público"],
            ["Body", '{ "correo": "ana@example.com" }'],
            ["Éxito", "200"],
            ["Mensaje", "Si el correo está registrado, te enviaremos instrucciones..."],
        ],
    )
    doc.add_paragraph("Respuesta (entorno de pruebas con RECOVERY_SIMULATION_EXPOSE=true):")
    code_block(
        doc,
        "{\n"
        '  "message": "Si el correo está registrado, te enviaremos instrucciones para restablecer la contraseña.",\n'
        '  "expiresInMinutes": 15,\n'
        '  "envio": {\n'
        '    "canal": "simulacion",\n'
        '    "asunto": "Restablece tu contraseña — OficiosYa",\n'
        '    "destinatario": "ana@example.com",\n'
        '    "nota": "No se despachó un correo real.",\n'
        '    "codigo": "847291",\n'
        '    "enlace": "http://localhost:5173/reset-password?token=...",\n'
        '    "vence": "2026-09-28T21:05:00.000Z"\n'
        "  }\n"
        "}",
    )
    doc.add_paragraph(
        "En producción (o con RECOVERY_SIMULATION_EXPOSE=false) no se devuelven codigo ni enlace; "
        "siguen visibles en la consola del backend como simulación de correo."
    )

    doc.add_heading("4. POST /auth/reset-password", 1)
    doc.add_paragraph("Hay dos formas equivalentes:")
    add_table(
        doc,
        ["Modo", "Body"],
        [
            ["Enlace", '{ "token": "<token>", "password": "nuevaClave123" }'],
            ["Código", '{ "correo": "ana@example.com", "codigo": "847291", "password": "nuevaClave123" }'],
        ],
    )
    add_table(
        doc,
        ["Campo", "Valor"],
        [
            ["Método / ruta", "POST /api/auth/reset-password"],
            ["Auth", "Público (el desafío sustituye al JWT)"],
            ["Éxito", "200 { message }"],
            ["401", "Código o enlace inválido, usado o expirado"],
            ["400", "Falta token/código o password < 8 caracteres"],
        ],
    )

    doc.add_heading("5. Cómo probarlo en Swagger", 1)
    doc.add_paragraph(
        "1. Abrir http://localhost:3000/api/docs\n"
        "2. POST /auth/forgot-password con un correo registrado\n"
        "3. Copiar envio.codigo o envio.enlace (o leer la consola del API: [HU-03] Simulación...)\n"
        "4. POST /auth/reset-password con token o correo+código y la nueva clave\n"
        "5. POST /auth/login con la clave nueva"
    )

    doc.add_heading("6. Tag Auth en OpenAPI", 1)
    doc.add_paragraph(
        "Los endpoints están en el tag Auth, esquemas ForgotPasswordDto, ResetPasswordDto, "
        "ForgotPasswordResponseDto y ResetPasswordResponseDto. La UI de /api/docs es la fuente viva."
    )

    doc.add_heading("7. Perfil y disponibilidad (IDC-21, HU-05 / HU-07)", 1)
    doc.add_paragraph(
        "GET/PUT /api/worker/profile gestionan tarifas, horarios y cobertura. "
        "PUT /api/worker/availability alterna Disponible/Ocupado. Requieren Bearer JWT."
    )
    add_table(
        doc,
        ["Ruta", "Notas"],
        [
            ["GET /api/worker/profile", "Borrador con existe=false si aún no hay perfil"],
            ["PUT /api/worker/profile", "Upsert. Tarifas/horarios en JSON de descripcion; cobertura en perfil_zona"],
            ["PUT /api/worker/availability", "Body opcional { disponibilidad }. Si no hay perfil, lo crea"],
        ],
    )

    doc.save(OUT)
    OUT_COPY.write_bytes(OUT.read_bytes())
    print(f"wrote {OUT}")
    print(f"wrote {OUT_COPY}")


if __name__ == "__main__":
    main()
